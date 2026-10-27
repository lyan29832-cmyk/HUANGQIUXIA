#!/usr/bin/env python3
"""Generate the restructured equipment classification and numbering SMP."""

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, Twips

OUT = "/workspace/生产设备与仪器分类、编号管理规程.docx"


def set_run_font(run, size=12, bold=False):
    run.bold = bold
    run.italic = False
    run.font.size = Pt(size)
    run.font.name = "宋体"
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rFonts.set(qn(attr), "宋体")


def set_paragraph_format(p, before=0, after=3, line=240, first_line=0, align="left"):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = line / 240
    pf.first_line_indent = Cm(first_line)
    pf.left_indent = Cm(0)
    align_map = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }
    p.alignment = align_map[align]


def add_text(p, text, size=12, bold=False):
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return run


def add_para(doc, text, size=12, bold=False, before=0, after=3, align="justify"):
    p = doc.add_paragraph()
    set_paragraph_format(p, before=before, after=after, align=align)
    add_text(p, text, size=size, bold=bold)
    return p


def add_heading_para(doc, text, level):
    size = 16 if level == 0 else (14 if level == 1 else 12)
    before = 0 if level == 0 else (10 if level == 1 else 6)
    after = 8 if level == 0 else 4
    align = "center" if level == 0 else "left"
    return add_para(doc, text, size=size, bold=True, before=before, after=after, align=align)


def shade_cell(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def set_cell_border(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn("w:tcBorders"))
    if tcBorders is None:
        tcBorders = OxmlElement("w:tcBorders")
        tcPr.append(tcBorders)
    for edge in ("top", "left", "bottom", "right"):
        el = tcBorders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            tcBorders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")


def set_cell_width(cell, width_cm):
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn("w:tcW"))
    if tcW is None:
        tcW = OxmlElement("w:tcW")
        tcPr.append(tcW)
    tcW.set(qn("w:w"), str(int(width_cm * 567)))
    tcW.set(qn("w:type"), "dxa")


def fill_cell(cell, text, bold=False, size=10.5, center=False, fill=None):
    cell.text = ""
    p = cell.paragraphs[0]
    set_paragraph_format(p, before=1, after=1, line=216, align="center" if center else "left")
    add_text(p, text, size=size, bold=bold)
    set_cell_border(cell)
    if fill:
        shade_cell(cell, fill)
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    vAlign = tcPr.find(qn("w:vAlign"))
    if vAlign is None:
        vAlign = OxmlElement("w:vAlign")
        tcPr.append(vAlign)
    vAlign.set(qn("w:val"), "center")


def prevent_row_split(row):
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    trPr.append(cant)


def add_table(doc, headers, rows, widths, header=True):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    total = int(sum(widths) * 567)
    tblW.set(qn("w:w"), str(total))
    tblW.set(qn("w:type"), "dxa")

    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        fill_cell(cell, h, bold=True, size=10.5, center=True, fill="D9E2F3")
        set_cell_width(cell, widths[i])
    prevent_row_split(table.rows[0])
    if header:
        trPr = table.rows[0]._tr.get_or_add_trPr()
        tblHeader = OxmlElement("w:tblHeader")
        trPr.append(tblHeader)

    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            fill_cell(cell, value, bold=False, size=10.5, center=False, fill="FFFFFF")
            set_cell_width(cell, widths[c_idx])
        prevent_row_split(table.rows[r_idx + 1])
    return table


def add_caption(doc, text):
    return add_para(doc, text, size=10.5, bold=True, before=8, after=2, align="center")


def add_page_field(paragraph, instr):
    run = paragraph.add_run()
    set_run_font(run, size=9, bold=False)
    r = run._r
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr_el = OxmlElement("w:instrText")
    instr_el.set(qn("xml:space"), "preserve")
    instr_el.text = instr
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    r.append(fld1)
    r.append(instr_el)
    r.append(fld2)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(12)
    rPr = normal.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rFonts.set(qn(attr), "宋体")


def configure_section(section):
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.4)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(0.8)

    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.clear()
    set_paragraph_format(hp, before=0, after=0, line=216, align="center")
    add_text(hp, "生产设备分类与编号管理规程", size=9, bold=False)

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.clear()
    set_paragraph_format(fp, before=0, after=0, line=216, align="center")
    add_text(fp, "修订稿    第 ", size=9)
    add_page_field(fp, " PAGE ")
    add_text(fp, " 页 / 共 ", size=9)
    add_page_field(fp, " NUMPAGES ")
    add_text(fp, " 页", size=9)

    settings = section._sectPr
    # fields update on open is set at document level


def enable_update_fields(doc):
    settings = doc.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def build():
    doc = Document()
    configure_styles(doc)
    configure_section(doc.sections[0])
    enable_update_fields(doc)

    add_heading_para(doc, "生产设备分类与编号管理规程", 0)
    add_para(doc, "（修订稿）", size=12, bold=True, before=0, after=6, align="center")

    add_caption(doc, "文件信息")
    add_table(
        doc,
        ["项目", "内容"],
        [
            ["文件名称", "生产设备分类与编号管理规程"],
            ["文件编号", "按公司文件编号规则填写"],
            ["版本号", "按本次修订的版本号填写"],
            ["起草部门", "设备工程部"],
            ["审核", "按文件批准权限填写"],
            ["批准", "按文件批准权限填写"],
            ["生效日期", "批准后填写"],
            ["状态", "修订稿。批准前不替代现行受控文本。"],
        ],
        [3.6, 12.3],
    )

    add_heading_para(doc, "1  目的", 1)
    add_para(
        doc,
        "1.1  本规程规定生产设备的分类原则、分类后的管理要求和唯一编号规则，并规定仪器的编号规则。仪器只按本规程编号，不按本规程分类，也不按本规程实施使用、维护、确认和校准管理。",
    )

    add_heading_para(doc, "2  范围", 1)
    add_para(
        doc,
        "2.1  本规程的设备分类和第6.3款管理要求，适用于公司生产设备、独立辅助设备、配电及电力设备和办公及其他设备。这些设备同时适用本规程的编号规则。",
    )
    add_para(
        doc,
        "2.2  仪器只适用本规程第6.4款至第6.8款的编号规则，包括编号结构、代号、配号、变更、注销，以及台账中的编号信息。仪器的分类、校准、使用、维护和记录执行其适用的管理规程。",
    )
    add_para(
        doc,
        "2.3  质量控制部仪器的分类、校准和检验使用管理执行 SMP-QC-000-021《检验仪器管理规程》。该类仪器的编号执行本规程，并登记在公司设备台账中。",
    )
    add_para(
        doc,
        "2.4  设备的系统影响性评估和组件关键性评估执行 SMP-QA-000-037《系统影响性评估及组件关键性评估管理规程》。本规程根据该评估结论确定设备的 A 类或 B 类管理要求。",
    )

    add_heading_para(doc, "3  职责", 1)
    add_para(
        doc,
        "3.1  设备工程部负责人与 QA 监督本规程的实施。发现编号重复、设备管理类别与评估结论不一致，或现场标识与设备台账不一致时，应当督促责任部门纠正。",
    )
    add_para(
        doc,
        "3.2  设备工程部维护公司设备台账，按本规程配发编号、登记编号变更并冻结已注销编号，保证现行编号唯一。",
    )
    add_para(
        doc,
        "3.3  设备归属部门提供设备的使用地点、用途和系统影响性评估结论；指定 A 类设备负责人；组织本部门设备操作规程、清洁规程和维护规程的起草、修订、审核和培训；完成设备负责人交接。",
    )
    add_para(
        doc,
        "3.4  仪器所在部门应当提供仪器名称、型号、出厂编号、安装地点和编号所需的大类、部门信息，并使现场仪器编号与设备台账一致。仪器的分类、校准和使用管理不由本规程规定。质量控制部仪器的上述管理执行 SMP-QC-000-021。",
    )
    add_para(
        doc,
        "3.5  设备使用人员按现场编号和状态标识使用设备，并如实填写使用日志及其他适用记录。仪器使用人员使用本规程配发的编号；仪器的使用和记录执行其适用规程。",
    )

    add_heading_para(doc, "4  编制依据", 1)
    add_para(
        doc,
        "4.1  《药品生产质量管理规范（2010年修订）》（卫生部令第79号）第七十一条、第七十二条、第七十三条、第七十九条、第八十条、第八十一条、第八十二条、第八十三条、第八十六条、第八十七条和第八十八条。上述条款在本规程中用于设备。仪器的校准和使用管理执行其适用规程。",
    )
    add_para(
        doc,
        "4.2  本规程引用的企业内部制度见下表。引用文本修订时，活动按修订后的受控文本执行。",
    )
    add_caption(doc, "表1  引用的企业内部制度")
    add_table(
        doc,
        ["文件编号", "文件名称", "在本规程中的用途"],
        [
            ["SMP-QA-000-037", "系统影响性评估及组件关键性评估管理规程", "确定设备的直接影响系统、非直接影响系统及组件关键性"],
            ["SMP-QA-000-030", "变更控制管理制度", "改变管理类别、编号或代码表"],
            ["SMP-VP-000-004", "设施、设备确认管理规程", "确定确认对象并实施确认、再确认"],
            ["SMP-EQ-100-003", "设备预防性维护、维修管理规程", "预防性维护及其维修安排"],
            ["SMP-EQ-100-004", "设备设施故障报告与维修管理规程", "故障报告与维修"],
            ["SMP-QC-000-021", "检验仪器管理规程", "质量控制部仪器的分类、校准和检验使用管理；本规程只为其配号"],
        ],
        [3.8, 6.4, 5.7],
    )
    add_para(
        doc,
        "4.3  本规程中的“确认”对应 qualification，“预防性维护”对应 preventive maintenance，设备编号用于实现 unique identification。中国生产现场按第4.1款和表1执行。",
        before=6,
    )

    add_heading_para(doc, "5  定义", 1)
    definitions = [
        "5.1  设备归属部门：对设备的日常使用、操作规程和现场管理承担责任的部门。",
        "5.2  直接影响系统：按 SMP-QA-000-037 评估，运行或失效预期会对产品质量产生直接影响的系统。",
        "5.3  非直接影响系统：按 SMP-QA-000-037 评估，自身不直接决定产品质量，而用于支持直接影响系统的系统。",
        "5.4  A 类设备：系统影响性评估结论为直接影响系统的设备。",
        "5.5  B 类设备：系统影响性评估结论为非直接影响系统的设备。",
        "5.6  独立辅助设备：可以独立使用，或同时服务于一台以上主设备，并单独列入设备台账的辅助设备。其大类代号为 F。",
        "5.7  附属装置：固定服务于一台主设备，并随该主设备一并管理的装置。附属装置使用该主设备的二级编号。",
        "5.8  一级编号：由设备大类代号、部门代号和三位流水号依次连接形成的编号。",
        "5.9  二级编号：在主设备一级编号后加上连字符和两位附属装置序号形成的编号。",
        "5.10  设备台账：登记设备的现行编号、管理类别、位置和状态，以及仪器现行编号的受控清单。经批准的纸质台账与电子台账均用于追溯。",
        "5.11  使用日志：用于药品生产的设备记录，内容包括使用、清洁、维护和维修情况以及日期、时间、所生产药品的名称、规格和批号等。本规程所称使用记录与使用日志为同一类记录。仪器的使用和校准记录不由本规程规定。",
        "5.12  仪器：本规程中需要编号的检测、测量或记录器具，大类代号为 J。仪器只执行编号规则。",
        "5.13  SMP-QA-000-037 对直接影响系统、非直接影响系统另有书面定义时，设备评估活动以该规程的定义为准。本规程第5.2款和第5.3款用于把设备的评估结论对应到 A 类、B 类。",
    ]
    for item in definitions:
        add_para(doc, item)

    add_heading_para(doc, "6  内容", 1)
    add_heading_para(doc, "6.1  分类与编号的关系", 2)
    add_para(
        doc,
        "6.1.1  管理类别表示设备所属系统对产品质量的影响，只用于设备。编号表示设备或仪器的身份。设备的类别和编号同时登记在设备台账中。大类代号不代替 A 类或 B 类，A 类或 B 类不代替设备编号。",
    )
    add_para(
        doc,
        "6.1.2  纳入设备台账的每台设备、仪器只有一个现行编号。完整编号在公司内唯一。",
    )
    add_para(
        doc,
        "6.1.3  编号在该设备或仪器存续期间保持稳定。设备或仪器报废、调出，或经变更控制批准重新编号后，原编号冻结，不再授予其他设备或仪器。历史记录继续沿用记录形成时的编号。",
    )
    add_para(
        doc,
        "6.1.4  仪器不划分 A 类或 B 类，不填写本规程的管理类别，只领取并使用编号。",
    )

    add_heading_para(doc, "6.2  设备分类", 2)
    add_para(
        doc,
        "6.2.1  本节只适用于设备。设备投入使用前，归属部门按 SMP-QA-000-037 完成系统影响性评估。组件关键性评估按该规程执行，其结论用于设备确认和维护范围的确定。本规程不另设组件分级。",
    )
    add_para(doc, "6.2.2  系统影响性评估结论按表2归入管理类别。")
    add_caption(doc, "表2  管理类别")
    add_table(
        doc,
        ["管理类别", "评估结论", "典型示例"],
        [
            ["A 类", "直接影响系统", "压片机、铝塑包装机、制粒机、空调净化系统、压缩空气系统"],
            ["B 类", "非直接影响系统", "叉车、洗衣机、干衣机"],
        ],
        [2.6, 4.2, 9.1],
    )
    add_para(
        doc,
        "6.2.3  表2中的示例用于帮助理解类别。单台设备的管理类别以该设备的系统影响性评估结论为准。评估结论为直接影响系统的设备按 A 类管理；评估结论为非直接影响系统的设备按 B 类管理。",
        before=6,
    )
    add_para(
        doc,
        "6.2.4  评估结论既不是直接影响系统，也不是非直接影响系统时，归属部门按 SMP-QA-000-037 确定其管理方式，并按 SMP-QA-000-030 批准后再登记管理类别。",
    )
    add_para(
        doc,
        "6.2.5  管理类别登记后保持不变。需要改变类别时，归属部门重新进行系统影响性评估，并按 SMP-QA-000-030 批准。变更批准前，设备继续按原类别管理。",
    )
    add_para(
        doc,
        "6.2.6  仪器不按本节分类。质量控制部仪器的分类执行 SMP-QC-000-021。其他仪器的分类执行其适用的管理规程。",
    )

    add_heading_para(doc, "6.3  分类管理要求", 2)
    add_para(
        doc,
        "6.3.1  本节只适用于设备。仪器不执行本节，只执行第6.4款至第6.8款。",
    )
    add_para(
        doc,
        "6.3.2  归属部门在设备投入使用前完成与其管理类别相适应的规程批准，并对操作人员和维护人员进行培训。培训合格后，人员方可独立操作或维护该设备。",
    )
    add_para(
        doc,
        "6.3.3  用于药品生产的设备应当建立使用日志。使用日志至少包括使用、清洁、维护和维修情况以及日期、时间、所生产药品的名称、规格和批号等。直接用于某一批产品生产的设备，使用日志应当记录该药品的名称、规格和批号。多个批次共用的公用系统，使用日志应当记录运行、清洁、维护、维修以及日期和时间，并使批记录能够追溯到该系统相应时间的运行记录。",
    )
    add_para(
        doc,
        "6.3.4  不用于药品生产的 B 类设备，由归属部门保存与其实际使用、清洁和维护相适应的记录。",
    )
    add_para(
        doc,
        "6.3.5  生产设备的状态标识应当标明设备编号。有内容物时，还应当标明内容物的名称、规格和批号；无内容物时，应当标明清洁状态。不合格设备在移出生产区或质量控制区之前，应当有醒目的状态标识。",
    )
    add_para(
        doc,
        "6.3.6  A 类设备由归属部门分别组织操作规程、清洁规程和维护规程的起草、修订和审核，并对相关人员进行培训。A 类设备建立第6.3.3款规定的使用日志，如实记录设备运行情况。",
    )
    add_para(
        doc,
        "6.3.7  A 类设备由归属部门指定设备负责人。设备负责人离职或调岗时，应当与继任者完成交接。交接内容至少包括使用日志、维护保养记录、操作规程、维护规程、清洁规程、备品备件、设备运行情况及注意事项。交接记录由双方签署姓名和日期，并纳入设备档案。",
    )
    add_para(
        doc,
        "6.3.8  B 类设备由归属部门根据设备的结构复杂程度和设备说明书，组织操作规程、清洁规程和维护规程的起草、修订和审核。操作人员和维护人员经培训合格后方可上岗。",
    )
    add_para(
        doc,
        "6.3.9  结构简单、操作和维护要求明确的 B 类设备，可以将操作规程、清洁规程和维护规程合并为一份《操作维护清洁规程》。合并后的规程应当完整、清晰、可执行，并包含使用、清洁和维护要求。",
    )
    add_para(
        doc,
        "6.3.10  设备确认的对象、方案和实施执行 SMP-VP-000-004。系统影响性评估结论作为确定确认范围的输入。已完成确认的生产设备，应当在确认的参数范围内使用。设备经改造或重大维修后，应当完成再确认并符合要求，方可恢复用于生产。",
    )
    add_para(
        doc,
        "6.3.11  预防性维护执行 SMP-EQ-100-003。故障报告与维修执行 SMP-EQ-100-004。设备的维护和维修记录使用本规程规定的设备编号。",
    )

    add_heading_para(doc, "6.4  编号规则", 2)
    add_para(doc, "6.4.1  一级编号由以下三段依次连接，段与段之间不插入空格或其他分隔符：")
    add_para(doc, "（1）设备大类代号；")
    add_para(doc, "（2）部门代号；")
    add_para(doc, "（3）三位流水号。")
    add_para(
        doc,
        "S1、S2 各占两个字符，其余大类代号占一个字符。部门代号占一个字符。一级编号因此为五个字符或六个字符。",
    )
    add_para(
        doc,
        "6.4.2  二级编号由主设备的一级编号、半角连字符“-”和两位附属装置序号依次连接组成。",
    )
    add_para(
        doc,
        "6.4.3  固定服务于一台主设备、并随该主设备一并管理的附属装置，使用二级编号。可以独立使用，或同时服务于多台主设备的辅助设备，使用大类 F 的一级编号。",
    )
    add_para(
        doc,
        "6.4.4  流水号在同一大类代号与同一部门代号的组合内从 001 起按发放顺序编制。附属装置序号在同一主设备下从 01 起按发放顺序编制。已配发的号码取消使用后，不再配发给其他设备；台账中可以保留空号。",
    )
    add_para(
        doc,
        "6.4.5  编号中的字母使用大写。台账、现场标识、记录和电子系统统一使用 S1、S2。",
    )

    add_heading_para(doc, "6.5  代号", 2)
    add_para(doc, "6.5.1  设备大类代号按表3使用。")
    add_caption(doc, "表3  设备大类代号")
    add_table(
        doc,
        ["代号", "含义", "代号", "含义"],
        [
            ["S1", "原料药车间生产设备", "S2", "口服固体制剂车间生产设备"],
            ["F", "辅助设备（独立辅助设备）", "P", "配电及电力设备"],
            ["Q", "办公及其他设备", "J", "检测仪器"],
        ],
        [3.2, 4.75, 3.2, 4.75],
    )
    add_para(doc, "6.5.2  部门代号按表4使用。", before=6)
    add_caption(doc, "表4  部门代号")
    add_table(
        doc,
        ["代号", "车间或部门", "代号", "车间或部门"],
        [
            ["A", "办公室", "B", "301车间、302车间"],
            ["C", "401车间", "D", "设备工程部"],
            ["E", "质量控制部", "F", "物料部"],
            ["G", "质量保证部", "H", "其他"],
        ],
        [3.2, 4.75, 3.2, 4.75],
    )
    add_para(
        doc,
        "6.5.3  原料药车间的生产设备使用大类 S1，口服固体制剂车间的生产设备使用大类 S2。部门代号按设备实际所在车间或部门从表4选取。设备位于 301 车间或 302 车间时，部门代号为 B；位于 401 车间时，部门代号为 C。",
        before=6,
    )
    add_para(
        doc,
        "6.5.4  部门代号 B 同时代表 301 车间和 302 车间。这两个车间在同一大类下共用一个流水号序列，按第6.4.4款连续配号。",
    )
    add_para(
        doc,
        "6.5.5  大类代号 F 与部门代号 F 字母相同，按“大类在前、部门在后”识读。例如 FF001 表示物料部的一台独立辅助设备，其流水号为 001。",
    )
    add_para(
        doc,
        "6.5.6  大类 Q 表示办公及其他设备。部门代号 H 表示表4所列部门以外的部门。使用 H 时，设备台账应当写明实际部门名称。新增部门代号或大类代号前，先按 SMP-QA-000-030 修订本规程的代码表，再按新代号配号。",
    )
    add_para(
        doc,
        "6.5.7  空调净化系统、压缩空气系统等公用系统的大类代号，按该系统在台账中的归属选用 S1、S2 或其他已列代号。其管理类别按该系统的系统影响性评估结论登记为 A 类或 B 类。",
    )
    add_para(
        doc,
        "6.5.8  检测仪器使用大类 J，部门代号按其实际归属选取。质量控制部的检测仪器，部门代号为 E。",
    )

    add_heading_para(doc, "6.6  编号的建立、使用、变更和注销", 2)
    add_para(
        doc,
        "6.6.1  设备归属部门在设备建账时，应当向设备工程部提供设备名称、型号、出厂编号、使用地点、建议的大类代号和部门代号，以及系统影响性评估结论或评估编号。附属装置同时提供主设备的现行编号。仪器所在部门只提供仪器名称、型号、出厂编号、安装地点和建议的大类代号、部门代号，不提供 A 类或 B 类结论。",
    )
    add_para(
        doc,
        "6.6.2  设备工程部应当核对代号和已发放的流水号，并将编号写入设备台账。编号自写入台账之日起启用。",
    )
    add_para(
        doc,
        "6.6.3  设备的现场标识、使用日志、维护维修记录和确认记录应当使用该设备的现行编号。仪器的现场编号使用本规程配发的编号。仪器的使用、校准和维护记录执行其适用规程，并引用同一编号。",
    )
    add_para(
        doc,
        "6.6.4  大类代号、部门代号或主从关系发生变化，因而需要改变编号时，按 SMP-QA-000-030 办理。变更批准后，设备工程部配发新编号。设备台账保留原编号、新编号、变更原因和生效日期。变更前已经形成的记录保持原编号。",
    )
    add_para(
        doc,
        "6.6.5  安装地点发生变化，而大类代号、部门代号和主从关系保持不变时，继续使用原编号，并在设备台账中更新安装地点。301 车间与 302 车间之间调动设备时，部门代号仍为 B，编号不变。",
    )
    add_para(
        doc,
        "6.6.6  设备或仪器报废、调出或不再纳入本公司管理时，设备工程部在台账中注销其在用状态，并按第6.1.3款冻结编号。",
    )
    add_para(
        doc,
        "6.6.7  本规程生效前已经发放、并且仍在使用的编号继续有效。发现历史编号与第6.4款或第6.5.4款不一致时，设备工程部在台账中登记差异；需要改号时按第6.6.4款办理。改号批准前继续使用原编号。历史编号使用下标 S₁、S₂ 的，在更换现场标识或整理台账时统一写为 S1、S2，并在台账中保留两种写法的对应关系。这种统一写法不视为重新编号，不冻结原编号。",
    )

    add_heading_para(doc, "6.7  设备台账", 2)
    add_para(
        doc,
        "6.7.1  设备工程部在设备台账中登记设备的现行编号，并使设备条目至少能够追溯下列信息：曾用编号、设备名称、型号、出厂编号、归属部门、安装地点、大类代号、管理类别、系统影响性评估编号、主设备编号（仅附属装置填写）、状态和启用日期。",
    )
    add_para(
        doc,
        "6.7.2  仪器与设备使用同一套编号体系，并登记在同一设备台账中。仪器条目登记编号、名称、型号、出厂编号、所在部门、安装地点、大类代号、状态和启用日期。仪器条目不按本规程填写管理类别和系统影响性评估编号。仪器所在部门可以另维护仪器清单，清单中的编号与设备台账一致。",
    )

    add_heading_para(doc, "6.8  编号示例", 2)
    add_para(doc, "6.8.1  下列流水号用于演示编号方法，不代表相应设备必须使用该号码。")
    add_para(
        doc,
        "6.8.2  401 车间的一台全自动胶囊填充机，属于口服固体制剂车间生产设备。假设流水号为 009，则一级编号为 S2C009。其中 S2 为口服固体制剂车间生产设备，C 为 401 车间，009 为流水号。",
    )
    add_para(
        doc,
        "6.8.3  401 车间的一台压片机，假设一级编号为 S2C080。固定服务于该压片机的第 1 台吸尘器属于附属装置，二级编号为 S2C080-01。该吸尘器不再另编大类 F 的一级编号。",
    )
    add_para(
        doc,
        "6.8.4  质量控制部的一台检测仪器，假设流水号为 001。其编号为 JE001。其中 J 为检测仪器，E 为质量控制部，001 为流水号。该仪器只取得编号，不登记 A 类或 B 类。",
    )
    add_caption(doc, "表5  示例编号的组成")
    add_table(
        doc,
        ["设备", "大类", "部门", "流水号", "附属序号", "完整编号"],
        [
            ["401车间全自动胶囊填充机", "S2", "C", "009", "—", "S2C009"],
            ["上述压片机", "S2", "C", "080", "—", "S2C080"],
            ["该压片机第1台专用吸尘器", "同主设备", "同主设备", "同主设备", "01", "S2C080-01"],
        ],
        [4.4, 2.3, 2.5, 2.4, 2.0, 2.3],
    )

    add_heading_para(doc, "7  注意事项", 1)
    add_para(
        doc,
        "7.1  确定设备管理类别时，以该设备的系统影响性评估结论为准，并以表2帮助理解 A 类、B 类所对应的结论。仪器不按表2分类。",
    )
    add_para(
        doc,
        "7.2  配号时按大类代号、部门代号、流水号的顺序连接。301 车间和 302 车间在同一大类下继续使用该序列尚未占用的下一流水号。",
    )
    add_para(
        doc,
        "7.3  附属装置的现场标识、使用日志和维修记录使用完整二级编号。该装置改为独立辅助设备时，按第6.6.4款配发新的一级编号，并冻结原二级编号。",
    )
    add_para(
        doc,
        "7.4  发现现场编号与设备台账不一致时，使用部门通知设备工程部。核对完成前，不把该不一致编号上的使用记录、维修记录或确认状态转移到另一台设备。设备工程部核对后更正现场标识或台账。已经形成的历史记录保持原样，并在台账中说明差异。",
    )
    add_para(
        doc,
        "7.5  仪器只执行本规程的编号。质量控制部仪器的分类、校准和检验使用管理执行 SMP-QC-000-021。其他仪器的分类、校准、使用和维护执行其适用规程。",
    )

    add_heading_para(doc, "8  环境、健康、安全（EHS）要求", 1)
    add_para(
        doc,
        "8.1  设备分类和编号活动不另行设置环境、职业健康或安全限度。",
    )
    add_para(
        doc,
        "8.2  设备安装、清洁、维修和公用系统操作中的职业安全与能源隔离，执行相应安全操作规程，以及 SMP-EQ-100-003、SMP-EQ-100-004。",
    )

    add_heading_para(doc, "9  派生记录", 1)
    add_heading_para(doc, "9.1  纸质受控记录", 2)
    add_para(
        doc,
        "9.1.1  本规程使用公司现有设备台账，不另建台账表单。台账信息满足第6.7.1款。",
    )
    add_para(
        doc,
        "9.1.2  A 类设备负责人交接记录纳入设备档案。公司已有交接单的，使用该交接单。尚未指定表单的，由文件管理部门指定记录表后实施第6.3.7款的交接记录。",
    )
    add_heading_para(doc, "9.2  电子模板受控记录", 2)
    add_para(
        doc,
        "9.2.1  设备台账采用电子系统时，该系统即本规程的设备台账。系统保证现行编号唯一，并保留管理类别变更、重新编号和注销的痕迹。本规程不另建电子记录模板。",
    )
    add_para(
        doc,
        "9.2.2  设备的使用日志、维护维修记录、确认记录、培训记录和变更记录，分别按各自的受控规程保存，其中的设备编号与设备台账的现行编号一致。仪器的使用、校准和维护记录按其适用规程保存，并引用本规程配发的编号。变更前的记录保留当时的编号。",
    )

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
