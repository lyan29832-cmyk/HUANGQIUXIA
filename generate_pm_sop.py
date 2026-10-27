#!/usr/bin/env python3
"""Generate the restructured equipment preventive-maintenance SMP."""

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, Twips

OUT = "/workspace/设备预防性维护保养规程.docx"


def set_run_font(run, size=12, bold=False, italic=False, latin="宋体"):
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = latin
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), latin)
    rFonts.set(qn("w:hAnsi"), latin)
    rFonts.set(qn("w:cs"), latin)
    rFonts.set(qn("w:eastAsia"), "宋体")


def set_paragraph_format(p, before=0, after=2, line=240, align="justify"):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = line / 240
    pf.first_line_indent = Cm(0)
    pf.left_indent = Cm(0)
    align_map = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }
    p.alignment = align_map[align]


def add_cn(doc, text, size=12, bold=False, before=0, after=1, align="justify"):
    p = doc.add_paragraph()
    set_paragraph_format(p, before=before, after=after, align=align)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, latin="宋体")
    return p


def add_en(doc, segments, size=12, before=0, after=8, align="justify"):
    """segments: str, or list of (text, italic) tuples."""
    if isinstance(segments, str):
        segments = [(segments, False)]
    p = doc.add_paragraph()
    set_paragraph_format(p, before=before, after=after, align=align)
    for text, italic in segments:
        run = p.add_run(text)
        set_run_font(run, size=size, bold=False, italic=italic, latin="Arial")
    return p


def add_heading(doc, cn, en, level):
    size = 16 if level == 0 else (14 if level == 1 else 12)
    before = 0 if level == 0 else (12 if level == 1 else 8)
    after = 6 if level == 0 else 2
    align = "center" if level == 0 else "left"
    p = doc.add_paragraph()
    set_paragraph_format(p, before=before, after=after, align=align)
    run = p.add_run(cn)
    set_run_font(run, size=size, bold=True, latin="宋体")
    run = p.add_run("  " + en)
    set_run_font(run, size=size, bold=True, latin="Arial")
    return p


def shade_cell(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def set_cell_border(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn("w:tcBorders"))
    if tcBorders is None:
        tcBorders = OxmlElement("w:tcBorders")
        tcPr.append(tcBorders)
    for edge in ("top", "left", "bottom", "right"):
        el = tcBorders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            tcBorders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")


def set_cell_width(cell, width_cm):
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn("w:tcW"))
    if tcW is None:
        tcW = OxmlElement("w:tcW")
        tcPr.append(tcW)
    tcW.set(qn("w:w"), str(int(width_cm * 567)))
    tcW.set(qn("w:type"), "dxa")


def fill_cell(cell, pieces, bold=False, size=10.5, fill=None):
    cell.text = ""
    p = cell.paragraphs[0]
    set_paragraph_format(p, before=2, after=2, line=216, align="left")
    if isinstance(pieces, str):
        pieces = [("cn", pieces, False)]
    for kind, text, italic in pieces:
        run = p.add_run(text)
        if kind == "cn":
            set_run_font(run, size=size, bold=bold, italic=False, latin="宋体")
        else:
            set_run_font(run, size=size, bold=bold, italic=italic, latin="Arial")
    set_cell_border(cell)
    if fill:
        shade_cell(cell, fill)
    tcPr = cell._tc.get_or_add_tcPr()
    vAlign = tcPr.find(qn("w:vAlign"))
    if vAlign is None:
        vAlign = OxmlElement("w:vAlign")
        tcPr.append(vAlign)
    vAlign.set(qn("w:val"), "center")


def prevent_row_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))


def add_ref_table(doc, rows):
    headers = ["编号 / Number", "文件 / Document"]
    widths = [4.2, 12.3]
    table = doc.add_table(rows=1 + len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl = table._tbl
    tblPr = tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(int(sum(widths) * 567)))
    tblW.set(qn("w:type"), "dxa")
    for i, h in enumerate(headers):
        fill_cell(table.rows[0].cells[i], h, bold=True, fill="D9E2F3")
        set_cell_width(table.rows[0].cells[i], widths[i])
    prevent_row_split(table.rows[0])
    trPr = table.rows[0]._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:tblHeader"))
    for r_idx, (num, cn, en) in enumerate(rows):
        fill_cell(table.rows[r_idx + 1].cells[0], num, bold=False)
        set_cell_width(table.rows[r_idx + 1].cells[0], widths[0])
        fill_cell(
            table.rows[r_idx + 1].cells[1],
            [("cn", cn + "\n", False), ("en", en, True)],
        )
        set_cell_width(table.rows[r_idx + 1].cells[1], widths[1])
        prevent_row_split(table.rows[r_idx + 1])
    return table


def add_page_number(paragraph):
    def fld(instr):
        r1 = paragraph.add_run()
        set_run_font(r1, size=9, latin="Arial")
        b1 = OxmlElement("w:fldChar")
        b1.set(qn("w:fldCharType"), "begin")
        r1._r.append(b1)
        r2 = paragraph.add_run()
        set_run_font(r2, size=9, latin="Arial")
        it = OxmlElement("w:instrText")
        it.set(qn("xml:space"), "preserve")
        it.text = instr
        r2._r.append(it)
        r3 = paragraph.add_run()
        set_run_font(r3, size=9, latin="Arial")
        b2 = OxmlElement("w:fldChar")
        b2.set(qn("w:fldCharType"), "end")
        r3._r.append(b2)

    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, latin="宋体")
    fld(" PAGE ")
    run = paragraph.add_run(" 页 / 共 ")
    set_run_font(run, size=9, latin="宋体")
    fld(" NUMPAGES ")
    run = paragraph.add_run(" 页")
    set_run_font(run, size=9, latin="宋体")


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.6)

    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(12)
    rPr = normal.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rFonts.set(qn(attr), "宋体")

    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    set_paragraph_format(hp, before=0, after=0, align="center")
    run = hp.add_run("设备预防性维护保养规程")
    set_run_font(run, size=9, latin="宋体")
    run = hp.add_run("  /  Equipment Preventive Maintenance Procedure")
    set_run_font(run, size=9, latin="Arial")

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    set_paragraph_format(fp, before=0, after=0, align="center")
    add_page_number(fp)

    add_heading(doc, "设备预防性维护保养规程", "Equipment Preventive Maintenance Procedure", 0)

    add_heading(doc, "1. 目的", "Purpose", 1)
    add_cn(doc, "制定本规程，旨在规范设备预防性维护保养工作，提升设备性能，降低故障率，确保生产活动持续符合药品生产质量管理规范要求。")
    add_en(doc, "This SMP is established to standardize preventive maintenance of equipment, improve equipment performance, reduce the failure rate, and ensure that production activities remain continuously compliant with GMP requirements.")

    add_heading(doc, "2. 范围", "Scope", 1)
    add_cn(doc, "2.1  本规程适用于设备的定期维护保养。")
    add_en(doc, "2.1  This SMP applies to the periodic maintenance of equipment.")
    add_cn(doc, "2.2  本规程对日常维护保养的评估、周期、记录和实施已作规定的，执行相应条款。")
    add_en(doc, "2.2  Where this SMP specifies the assessment, interval, records, or performance of routine maintenance, the corresponding clauses apply.")

    add_heading(doc, "3. 职责", "Responsibilities", 1)
    add_cn(doc, "3.1  设备工程部和生产部负责执行本规程。")
    add_en(doc, "3.1  The Equipment Engineering Department and the Production Department are responsible for implementing this SMP.")
    add_cn(doc, "3.2  质量保证部负责人和设备工程部负责人负责监督本规程的执行。")
    add_en(doc, "3.2  The head of the Quality Assurance Department and the head of the Equipment Engineering Department are responsible for overseeing the implementation of this SMP.")

    add_heading(doc, "4. 编制依据", "References", 1)
    add_cn(doc, "4.1  《药品生产质量管理规范》（现行版）。")
    add_en(doc, "4.1  Good Manufacturing Practice for Pharmaceutical Products (current version).")
    add_cn(doc, "4.2  《药品GMP指南——厂房设施与设备》（现行版）。")
    add_en(
        doc,
        [
            ("4.2  ", False),
            ("Pharmaceutical GMP Guide—Facilities and Equipment", True),
            (" (current version).", False),
        ],
    )
    add_cn(doc, "4.3  本规程引用下列已有文件，不另建文件：", before=4)
    add_en(doc, "4.3  This SMP cites the following existing documents and does not create new documents:", after=4)
    add_ref_table(
        doc,
        [
            ("SMP-DM-000-009", "《文件发放管理规程》", "Document Issuance Management Procedure"),
            ("SMP-EQ-100-004", "《设备设施故障报告与维修管理规程》", "Equipment and Facility Failure Reporting and Repair Management Procedure"),
            ("SMP-VP-000-003", "《验证管理规程》", "Validation Management Procedure"),
            ("SMP-QA-000-024", "《产品质量回顾分析管理规程》", "Product Quality Review and Analysis Management Procedure"),
        ],
    )

    add_heading(doc, "5. 定义", "Definitions", 1)
    add_cn(doc, "5.1  预防性维护保养：为最大限度减少对正常使用的影响、延长维护对象的使用寿命，按照事先制定的维护计划实施的事前维护活动，包括日常维护保养和定期预防性维护保养。")
    add_en(doc, "5.1  Preventive maintenance: maintenance activities performed in advance, in accordance with a predefined maintenance plan, in order to minimize to the greatest extent practicable the impact on normal use and to extend the service life of the item being maintained. Preventive maintenance includes routine maintenance and periodic preventive maintenance.")

    add_heading(doc, "6. 预防性维护保养", "Preventive Maintenance", 1)

    add_heading(doc, "6.1  评估", "Assessment", 2)
    add_cn(doc, "6.1.1  设备安装就位后，设备工程部和使用部门填写REC-EQ-100-036《设备预防性维护保养评估表》，确定是否需要实施预防性维护保养，并明确相应的维护保养周期。")
    add_en(
        doc,
        [
            ("6.1.1  After the equipment has been installed and positioned, the Equipment Engineering Department and the user department shall complete REC-EQ-100-036, ", False),
            ("Equipment Preventive Maintenance Assessment Form", True),
            (", to determine whether preventive maintenance is required and to define the corresponding maintenance interval.", False),
        ],
    )
    add_cn(doc, "6.1.2  日常维护保养由使用部门结合设备厂家建议、设备用途及工艺要求进行综合评估，确定是否需要实施。日常维护保养周期由使用部门根据实际使用情况确定。")
    add_en(doc, "6.1.2  The user department shall assess routine maintenance based on the equipment manufacturer's recommendations, the intended use of the equipment, and the process requirements, and shall determine whether routine maintenance is required. The user department shall set the routine-maintenance interval according to actual use.")
    add_cn(doc, "6.1.3  定期预防性维护保养由设备工程部根据设备厂家建议、REC-QA-500-009《系统影响性评估表》、设备分类、设备用途及工艺要求进行综合评估，确定是否需要实施。定期预防性维护保养周期由设备工程部根据设备用途、经验、设备回顾报告、风险分析、供应商的维护建议和实际使用情况制定。相同设备因用途不同，可以设定不同的维护周期。")
    add_en(
        doc,
        [
            ("6.1.3  The Equipment Engineering Department shall assess periodic preventive maintenance based on the equipment manufacturer's recommendations, REC-QA-500-009, ", False),
            ("System Impact Assessment Form", True),
            (", the equipment classification, the intended use of the equipment, and the process requirements, and shall determine whether periodic preventive maintenance is required. The Equipment Engineering Department shall set the interval according to the intended use of the equipment, experience, equipment review reports, risk analysis, the supplier's maintenance recommendations, and actual use. The same type of equipment may be assigned different intervals where the intended use differs.", False),
        ],
    )
    add_cn(doc, "6.1.4  定期预防性维护保养按下列类别执行：")
    add_en(doc, "6.1.4  Periodic preventive maintenance shall follow the categories below:", after=2)
    add_cn(doc, "（1）A类：通常情况下需要实施定期预防性维护保养。需校验的仪器设备，以及设备结构和电气控制相对简单的设备，可以不实施定期预防性维护保养。")
    add_en(doc, "(1) Category A: periodic preventive maintenance is generally required. Instruments and equipment subject to calibration, and equipment with a relatively simple mechanical structure and electrical control system, need not undergo periodic preventive maintenance.", after=2)
    add_cn(doc, "（2）B类：通常情况下不需要实施定期预防性维护保养。设备结构和电气控制相对复杂的设备，需要实施定期预防性维护保养。")
    add_en(doc, "(2) Category B: periodic preventive maintenance is generally not required. Equipment with a relatively complex mechanical structure and electrical control system shall undergo periodic preventive maintenance.")

    add_heading(doc, "6.2  维护保养项目", "Maintenance Activities", 2)
    add_cn(doc, "6.2.1  设备维护保养项目应当根据设备操作说明书、组件关键性评估等相关资料确定。")
    add_en(doc, "6.2.1  Equipment maintenance activities shall be determined from the equipment operating manual, component criticality assessments, and other relevant documentation.")
    add_cn(doc, "6.2.2  结合设备实际运行情况和回顾分析结果，可以适当调整维护保养项目，包括增加或减少项目。")
    add_en(doc, "6.2.2  Maintenance activities may be adjusted, including by adding or removing activities, based on the actual operating condition of the equipment and the results of review and analysis.")
    add_cn(doc, "6.2.3  维护保养项目不限于设备资料中的要求。凡可能影响设备安全、稳定运行的内容，均应当纳入维护保养范围。")
    add_en(doc, "6.2.3  Maintenance activities are not limited to those specified in the equipment documentation. Any activity that may affect the safe and stable operation of the equipment shall be included in the scope of maintenance.")

    add_heading(doc, "6.3  维护保养文件", "Maintenance Documentation", 2)
    add_cn(doc, "6.3.1  设备工程部根据设备的维护保养周期和预防性维护保养项目，制定相应的预防性维护保养操作规程和记录。日常维护保养记录由设备使用部门根据日常维护保养项目制定。")
    add_en(doc, "6.3.1  The Equipment Engineering Department shall prepare the preventive-maintenance procedures and records according to the maintenance interval and the preventive-maintenance activities. The user department shall prepare the routine-maintenance records according to the routine-maintenance activities.")

    add_heading(doc, "6.4  维护保养计划", "Maintenance Plan", 2)
    add_cn(doc, "6.4.1  预防性维护保养计划由设备工程部制定。计划内容通常包括设备名称、设备编号和维护时间。")
    add_en(doc, "6.4.1  The Equipment Engineering Department shall prepare the preventive maintenance plan. The plan generally includes the equipment name, equipment number, and maintenance time.")

    add_cn(doc, "6.4.2  年度预防性维护保养计划", bold=True, before=6, after=2)
    add_en(doc, "6.4.2  Annual preventive maintenance plan", after=4)
    add_cn(doc, "6.4.2.1  设备工程部根据各类设备的具体维护内容，以及各维护项目的时间和期限或周期，于每年12月底前制定次年的REC-EQ-100-007《主要设备预防性维护保养年度计划表》。该计划经设备工程部负责人审核，经生产管理负责人和质量管理负责人批准后，按照SMP-DM-000-009《文件发放管理规程》发放。")
    add_en(
        doc,
        [
            ("6.4.2.1  Based on the specific maintenance content of each type of equipment and the timing and interval of each maintenance activity, the Equipment Engineering Department shall prepare the next year's REC-EQ-100-007, ", False),
            ("Annual Preventive Maintenance Plan for Major Equipment", True),
            (", by the end of December. The head of the Equipment Engineering Department shall review the plan. The head of Production Management and the head of Quality Management shall approve it. The plan shall then be issued in accordance with SMP-DM-000-009, ", False),
            ("Document Issuance Management Procedure", True),
            (".", False),
        ],
    )
    add_cn(doc, "6.4.2.2  制定年度维护保养计划时，应当遵循集中维护保养原则。在不超过规定维护周期的前提下，统计下一年度所有相关生产设备在半年度、年度等周期内的维护保养项目及对应月份，尽量将维护任务集中安排在同一个适宜的月份实施，以提高维护效率，减少对生产的影响。")
    add_en(doc, "6.4.2.2  The annual maintenance plan shall follow the principle of grouping maintenance activities. Provided the specified maintenance interval is not exceeded, the maintenance activities due in the following year for all relevant production equipment at semi-annual, annual, and other stated intervals, together with the corresponding months, shall be identified and, to the extent practicable, scheduled in one suitable month. The purpose is to improve maintenance efficiency and reduce the impact on production.")

    add_cn(doc, "6.4.3  月度预防性维护保养计划", bold=True, before=6, after=2)
    add_en(doc, "6.4.3  Monthly preventive maintenance plan", after=4)
    add_cn(doc, "6.4.3.1  设备工程部根据《主要设备预防性维护保养年度计划表》，按自然月制定REC-EQ-100-006《月度设备维护保养计划》，并预估各维护项目所需工时。每月25日前，将下月计划以邮件发送至相关部门审核，确认维护时间与生产安排不冲突。审核无误后，由制表人和复核人签字，并按照SMP-DM-000-009《文件发放管理规程》发放。")
    add_en(
        doc,
        [
            ("6.4.3.1  The Equipment Engineering Department shall prepare REC-EQ-100-006, ", False),
            ("Monthly Equipment Maintenance Plan", True),
            (", for each calendar month from the ", False),
            ("Annual Preventive Maintenance Plan for Major Equipment", True),
            (", and shall estimate the working time required for each maintenance activity. By the 25th of each month, the plan for the following month shall be sent by email to the relevant departments for review, to confirm that the maintenance time does not conflict with the production schedule. After the review identifies no discrepancy, the preparer and the reviewer shall sign the plan. The plan shall be issued in accordance with SMP-DM-000-009, ", False),
            ("Document Issuance Management Procedure", True),
            (".", False),
        ],
    )
    add_cn(doc, "6.4.3.2  相关部门审核月度计划时，应当结合生产安排确定预防性维护保养的实施时间，确保维护活动不影响正常生产。")
    add_en(doc, "6.4.3.2  When reviewing the monthly plan, the relevant departments shall take the production schedule into account and determine the time for preventive maintenance, so that the maintenance does not affect normal production.")
    add_cn(doc, "6.4.3.3  编制月度计划时，排期以该设备上次实际维护日期为基准，可以适当提前或推迟，偏移不超过10天。设备工程部应当按照当月生效的计划实施维护保养。")
    add_en(doc, "6.4.3.3  When the monthly plan is prepared, the date shall be based on the equipment's last actual maintenance date. The date may be brought forward or postponed, provided the shift does not exceed 10 days. The Equipment Engineering Department shall carry out maintenance in accordance with the plan effective for that month.")
    add_cn(doc, "6.4.3.4  维护计划需要变更时，设备工程部应当在原计划文件上注明变更原因和变更内容，并及时通知相关部门。生产部调整生产计划时，应当至少提前2天通知设备工程部。")
    add_en(doc, "6.4.3.4  When the maintenance plan needs to be changed, the Equipment Engineering Department shall record the reason and content of the change on the original plan and notify the relevant departments promptly. When the Production Department changes the production schedule, it shall notify the Equipment Engineering Department at least 2 days in advance.")
    add_cn(doc, "6.4.3.5  原料药车间组合式空调机组（FD022）、原料药车间非防爆区空调系统（S1B033）和口服固体制剂车间空调系统（S2C013）的预防性维护保养，应当结合每月洁净区空间臭氧消毒计划，安排在臭氧消毒当日、臭氧消毒操作开始前实施。")
    add_en(doc, "6.4.3.5  Preventive maintenance of the combined air-handling unit in the API workshop (FD022), the HVAC system for the non-explosion-proof area of the API workshop (S1B033), and the HVAC system of the oral solid dosage workshop (S2C013) shall be arranged together with the monthly ozone-disinfection plan for the cleanroom space. The maintenance shall be performed on the day of ozone disinfection and before ozone disinfection starts.")
    add_cn(doc, "6.4.4  新增设备或设备发生转移时，设备工程部应当根据相应的预防性维护保养标准操作规程，及时更新相关维护保养计划。")
    add_en(doc, "6.4.4  When equipment is added or relocated, the Equipment Engineering Department shall update the relevant maintenance plan promptly, in accordance with the applicable preventive-maintenance standard operating procedure.")

    add_heading(doc, "6.5  实施", "Performance", 2)
    add_cn(doc, "6.5.1  一般要求", bold=True, before=4, after=2)
    add_en(doc, "6.5.1  General requirements", after=4)
    add_cn(doc, "6.5.1.1  日常维护保养由设备使用部门的操作人员负责执行。")
    add_en(doc, "6.5.1.1  Routine maintenance shall be performed by operators of the user department.")
    add_cn(doc, "6.5.1.2  定期预防性维护保养（月度、季度、半年度、年度）由设备维修人员主导，设备操作人员协助实施。")
    add_en(doc, "6.5.1.2  Periodic preventive maintenance at monthly, quarterly, semi-annual, and annual intervals shall be led by equipment maintenance personnel, with equipment operators assisting.")
    add_cn(doc, "6.5.1.3  因生产活动或其他特殊情况不能按计划日期实施维护保养时，设备工程部可以在计划日期的基础上提前或推迟，偏移不超过10天。超过10天仍不能实施的，应当评估风险，并在维护记录中注明推迟原因和风险评估结论。")
    add_en(doc, "6.5.1.3  Where maintenance cannot be performed on the planned date because of production activities or other special circumstances, the Equipment Engineering Department may bring it forward or postpone it from the planned date, provided the shift does not exceed 10 days. Where it still cannot be performed after 10 days, a risk assessment shall be performed, and the reason for the postponement and the conclusion of the risk assessment shall be recorded in the maintenance record.")
    add_cn(doc, "6.5.1.4  涉及配电柜的生产设备进行预防性维护保养时，应当在设备维护人员的指导协助下，由生产人员和维护人员共同对配电柜进行卫生清洁和检查。")
    add_en(doc, "6.5.1.4  During preventive maintenance of production equipment that has an electrical distribution cabinet, production personnel and maintenance personnel shall jointly clean and inspect the cabinet under the guidance and assistance of equipment maintenance personnel.")
    add_cn(doc, "6.5.1.5  在洁净区安装压力表及连接管道时，密封用的生料带不得裸露于接口处。")
    add_en(doc, "6.5.1.5  When a pressure gauge and its connecting pipe are installed in a cleanroom, the PTFE thread-seal tape used for sealing shall not be exposed at the joint.")
    add_cn(doc, "6.5.1.6  完成预防性维护保养项目（例如更换制冷剂、机油或关键组件）后，必须对设备进行试运行。")
    add_en(doc, "6.5.1.6  After a preventive-maintenance activity is completed (for example, replacement of refrigerant, lubricating oil, or a critical component), the equipment shall undergo a trial run.")
    add_cn(doc, "6.5.1.7  维护保养完成后，由使用部门审核确认，并判断是否需要清洁或消毒。需要清洁或消毒的，由使用部门执行。")
    add_en(doc, "6.5.1.7  After maintenance is completed, the user department shall review and confirm completion and decide whether cleaning or disinfection is required. Where cleaning or disinfection is required, the user department shall perform it.")
    add_cn(doc, "6.5.1.8  预防性维护保养过程中发现设备故障，且所需工作不属于预防性维护保养项目的，按照SMP-EQ-100-004《设备设施故障报告与维修管理规程》执行。")
    add_en(
        doc,
        [
            ("6.5.1.8  Where an equipment failure is found during preventive maintenance and the required work is not a preventive-maintenance activity, SMP-EQ-100-004, ", False),
            ("Equipment and Facility Failure Reporting and Repair Management Procedure", True),
            (", applies.", False),
        ],
    )
    add_cn(doc, "6.5.1.9  关键工艺设备的维护保养应当在现场QA监督下进行。关键工艺设备是指用于关键生产工序、产生关键工艺参数或影响关键质量属性的生产设备。")
    add_en(doc, "6.5.1.9  Maintenance of critical process equipment shall be performed under the oversight of on-site QA. Critical process equipment means production equipment used in a critical manufacturing step that generates a critical process parameter or affects a critical quality attribute.")

    add_cn(doc, "6.5.2  关键设备润滑", bold=True, before=6, after=2)
    add_en(doc, "6.5.2  Lubrication of critical equipment", after=4)
    add_cn(doc, "设备工程部按照《关键设备/设施清单》和关键设备预防性维护保养文件中的润滑项目，建立REC-EQ-100-053《关键设备部件润滑台帐》。维护保养过程中完成部件润滑后，操作人员应当及时在该台帐中登记润滑信息。")
    add_en(
        doc,
        [
            ("Based on the List of Critical Equipment/Facilities and the lubrication activities in the preventive-maintenance documents for critical equipment, the Equipment Engineering Department shall establish REC-EQ-100-053, ", False),
            ("Lubrication Log for Critical Equipment Components", True),
            (". After a component is lubricated during maintenance, the operator shall record the lubrication information in the log promptly.", False),
        ],
    )

    add_cn(doc, "6.5.3  关键组件更换", bold=True, before=6, after=2)
    add_en(doc, "6.5.3  Replacement of critical components", after=4)
    add_cn(doc, "6.5.3.1  设备工程部按照《关键设备/设施清单》和关键设备预防性维护保养项目，建立并维护REC-EQ-100-042《关键设备部件需周期更换台帐》。")
    add_en(
        doc,
        [
            ("6.5.3.1  The Equipment Engineering Department shall establish and maintain REC-EQ-100-042, ", False),
            ("Periodic Replacement Log for Critical Equipment Components", True),
            (", based on the List of Critical Equipment/Facilities and the preventive-maintenance activities for critical equipment.", False),
        ],
    )
    add_cn(doc, "6.5.3.2  组件更换通常与预防性维护保养同步实施。维护前，维修人员应当查阅该台帐，确认是否有到期需更换的组件，并在规定期限内完成更换。更换的关键组件应当与原组件的规格、型号相同。")
    add_en(doc, "6.5.3.2  Component replacement is generally performed at the same time as preventive maintenance. Before maintenance, maintenance personnel shall review the log, confirm whether any component is due for replacement, and complete the replacement within the specified period. A replacement critical component shall have the same specification and model as the original component.")
    add_cn(doc, "6.5.3.3  关键设备组件更换后，应当按照第6.5.4款规定的调试与风险评估流程执行。")
    add_en(doc, "6.5.3.3  After a critical component is replaced, the commissioning and risk-assessment process in Section 6.5.4 shall be followed.")
    add_cn(doc, "6.5.3.4  组件更换完成后，应当及时更新该台帐。台帐于每年年底打印，经部门复核、审核后归档保存。")
    add_en(doc, "6.5.3.4  The log shall be updated promptly after component replacement. The log shall be printed at the end of each year and, after departmental check and review, archived.")

    add_cn(doc, "6.5.4  拆卸与复原", bold=True, before=6, after=2)
    add_en(doc, "6.5.4  Disassembly and restoration", after=4)
    add_cn(doc, "6.5.4.1  生产设备的预防性维护保养涉及拆卸与复原时，应当严格按照标准操作规程或既定方案执行。操作前，应当准备所需工具、备件和防护用品，并对设备状态进行标识；必要时，对拆下的零部件编号和标记。复原前，应当确认所有零部件符合技术要求，并核对拆卸时的标识。应当按照标准操作规程或图纸复原，确保装配正确。")
    add_en(doc, "6.5.4.1  Where preventive maintenance of production equipment involves disassembly and restoration, the work shall be performed strictly in accordance with the applicable standard operating procedure or the established plan. Before the work starts, the required tools, spare parts, and personal protective equipment shall be prepared, and the equipment status shall be identified. Where necessary, removed parts shall be numbered and marked. Before restoration, it shall be confirmed that every part meets the technical requirements, and the identification made at disassembly shall be checked. Restoration shall follow the standard operating procedure or the drawing, and the assembly shall be correct.")
    add_cn(doc, "6.5.4.2  关键生产设备拆卸与复原后，应当填写附件一《设备预防性维护保养后的调试与确认表》。复原后应当进行必要的调试与确认，确认设备功能正常、状态可控。")
    add_en(
        doc,
        [
            ("6.5.4.2  For critical production equipment, Appendix 1, ", False),
            ("Commissioning and Confirmation Form for Equipment Following Preventive Maintenance", True),
            (", shall be completed after disassembly and restoration. After restoration, the equipment shall be commissioned and confirmed as necessary, to confirm that its function is normal and its status is under control.", False),
        ],
    )
    add_cn(doc, "6.5.4.3  维保人员填写附件一。基本信息包括设备名称、设备编号、拆卸的组件、是否涉及关键组件更换，以及是否复原。")
    add_en(doc, "6.5.4.3  Maintenance personnel shall complete Appendix 1. The basic information comprises the equipment name, equipment number, components disassembled, whether replacement of a critical component is involved, and whether the equipment has been restored.")
    add_cn(doc, "6.5.4.4  拆卸复原后，由维保人员和设备使用部门人员共同调试。调试过程由现场QA监督。调试完成后，由维保人员、设备使用部门人员和现场QA签名确认。调试内容包括：")
    add_en(doc, "6.5.4.4  After disassembly and restoration, maintenance personnel and personnel of the user department shall commission the equipment together. On-site QA shall oversee the commissioning. After commissioning, maintenance personnel, personnel of the user department, and on-site QA shall sign. Commissioning includes the following:", after=2)
    checks = [
        (
            "（1）设备外观完好，无破损，无异常。",
            "(1) The appearance of the equipment is intact, with no damage and no abnormality.",
        ),
        (
            "（2）控制面板及按钮操作灵敏、正常，反馈信号准确。",
            "(2) The control panel and buttons operate responsively and normally, and the feedback signals are accurate.",
        ),
        (
            "（3）各连接紧固，润滑以及气路、水路连接正确，控制系统上电正常，运行指示灯有效。",
            "(3) Connections are secure; lubrication and the pneumatic and water connections are correct; the control system powers on normally; and the running indicator functions.",
        ),
        (
            "（4）按照设备标准操作程序运行设备。设备启动和停止功能正常，运行平稳，无异常噪声、振动或泄漏。",
            "(4) The equipment is operated in accordance with its standard operating procedure. Start and stop functions are normal, operation is stable, and there is no abnormal noise, vibration, or leakage.",
        ),
        (
            "（5）设备控制参数（如温度、压力、转速等）稳定并处于设定范围内，设备功能正常。",
            "(5) Equipment control parameters, such as temperature, pressure, and rotational speed, are stable and within the set range, and the equipment functions normally.",
        ),
    ]
    for cn, en in checks:
        add_cn(doc, cn)
        add_en(doc, en, after=2)
    add_cn(doc, "6.5.4.5  调试完成后，设备使用部门和验证QA应当评估拆卸与复原是否影响设备的负载运行状态和验证状态，并提交设备工程部负责人、设备使用部门负责人和质量保证部负责人审核签字。评估包括以下内容：", before=4)
    add_en(doc, "6.5.4.5  After commissioning, the user department and Validation QA shall assess whether disassembly and restoration have affected the loaded operating condition and the validation status of the equipment. The assessment shall be submitted to the head of the Equipment Engineering Department, the head of the user department, and the head of the Quality Assurance Department for review and signature. The assessment includes the following:", after=2)
    add_cn(doc, "（1）使用模拟物料或实际物料进行负载运行，确认设备在负载条件下的运行状态符合要求。")
    add_en(doc, "(1) The equipment is run under load with simulated material or actual material, to confirm that its operating condition under load meets the requirements.", after=2)
    add_cn(doc, "（2）关键组件更换已经影响设备验证状态的，应当按照SMP-VP-000-003《验证管理规程》进行确认，并参考《系统影响性及组件关键性评估报告》。更换关键组件后，应当填写《维护保养更换关键组件风险评估表》，按失效模式对所更换的组件进行风险评估，确定确认的范围和深度。")
    add_en(
        doc,
        [
            ("(2) Where replacement of a critical component has affected the validated state of the equipment, qualification shall be performed in accordance with SMP-VP-000-003, ", False),
            ("Validation Management Procedure", True),
            (", with reference to the System Impact and Component Criticality Assessment Report. After the critical component is replaced, the ", False),
            ("Risk Assessment Form for Replacement of Critical Components during Maintenance", True),
            (" shall be completed. The replaced component shall be assessed by failure mode to determine the scope and depth of qualification.", False),
        ],
    )
    add_cn(doc, "6.5.4.6  负载运行测试和确认完成后，应当在附件一“执行结果”栏中选择结果是否符合要求，将负载测试记录附于附件一之后，填写确认报告编号，并由设备使用部门和验证QA签字确认。")
    add_en(doc, "6.5.4.6  After the load run and confirmation are completed, the “execution result” column of Appendix 1 shall be completed to show whether the result meets the requirements. The load-test record shall be attached to Appendix 1, the confirmation-report number shall be entered, and the user department and Validation QA shall sign.")
    add_cn(doc, "6.5.4.7  附件一作为维护保养记录的附件，与维护保养记录一并归档保存。")
    add_en(doc, "6.5.4.7  Appendix 1 is an attachment to the maintenance record and shall be archived together with that record.")

    add_heading(doc, "6.6  回顾分析", "Retrospective Review", 2)
    add_cn(doc, "6.6.1  设备工程部按照SMP-QA-000-024《产品质量回顾分析管理规程》，对设备维护保养情况进行年度回顾。回顾内容包括下列各项：")
    add_en(
        doc,
        [
            ("6.6.1  The Equipment Engineering Department shall perform an annual review of equipment maintenance in accordance with SMP-QA-000-024, ", False),
            ("Product Quality Review and Analysis Management Procedure", True),
            (". The review includes the following:", False),
        ],
        after=2,
    )
    add_cn(doc, "（1）评估维护保养的实施效果，识别是否需要优化维护项目或调整周期。")
    add_en(doc, "(1) Evaluate the effect of the maintenance performed, and identify whether maintenance activities need to be optimized or intervals adjusted.", after=2)
    add_cn(doc, "（2）在《产品质量回顾分析报告》中提出维护项目或周期的调整建议，经批准后执行。")
    add_en(
        doc,
        [
            ("(2) Propose adjustments to maintenance activities or intervals in the ", False),
            ("Product Quality Review and Analysis Report", True),
            (", and implement them after approval.", False),
        ],
        after=2,
    )
    add_cn(doc, "（3）其他需要调整维护项目或周期的情形，应当先进行风险评估，经审核批准后方可实施。")
    add_en(doc, "(3) For any other case in which a maintenance activity or interval needs to be adjusted, a risk assessment shall be performed first. The adjustment may be implemented only after review and approval.", after=2)
    add_cn(doc, "（4）每年回顾完成后，设备工程部应当结合回顾情况，组织召开“关键生产设备及设施预防性维护保养策略”专题讨论会，持续优化维护策略。")
    add_en(doc, "(4) After the annual review is completed, the Equipment Engineering Department shall, based on the review, hold a meeting on the preventive maintenance strategy for critical production equipment and facilities, and continue to optimize the maintenance strategy.")

    add_heading(doc, "7. 注意事项", "Precautions", 1)
    add_cn(doc, "所提供原稿止于本章标题。本章正文未在本次整理中补写。")
    add_en(doc, "The source text provided ends at this heading. The body of this section has not been added in this revision.")

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
